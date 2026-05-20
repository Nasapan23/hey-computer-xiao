from __future__ import annotations

import re
from pathlib import Path

from docx import Document
from docx.shared import Pt


PROJECT_ROOT = Path(__file__).resolve().parents[1]
INPUT_MD = PROJECT_ROOT / "report" / "report1_visual_fingerprint_guard_report.md"
OUTPUT_DOCX = PROJECT_ROOT / "report" / "report1_visual_fingerprint_guard_report.docx"
OUTPUT_DOCX_ALT = PROJECT_ROOT / "report" / "report1_visual_fingerprint_guard_report_v2.docx"


def clean_inline_md(text: str) -> str:
    text = re.sub(r"`([^`]+)`", r"\1", text)
    text = re.sub(r"\*\*([^*]+)\*\*", r"\1", text)
    return text


def is_table_separator(line: str) -> bool:
    s = line.strip()
    if not (s.startswith("|") and s.endswith("|")):
        return False
    core = s.strip("|").strip()
    return bool(core) and all(part.strip().replace("-", "").replace(":", "") == "" for part in core.split("|"))


def parse_table_line(line: str) -> list[str]:
    parts = [p.strip() for p in line.strip().strip("|").split("|")]
    return [clean_inline_md(p) for p in parts]


def add_code_block(doc: Document, code_lines: list[str]) -> None:
    text = "\n".join(code_lines)
    p = doc.add_paragraph()
    run = p.add_run(text)
    run.font.name = "Consolas"
    run.font.size = Pt(9)


def main() -> None:
    if not INPUT_MD.exists():
        raise FileNotFoundError(f"Missing markdown file: {INPUT_MD}")

    lines = INPUT_MD.read_text(encoding="utf-8").splitlines()
    doc = Document()

    in_code = False
    code_lines: list[str] = []
    i = 0
    while i < len(lines):
        line = lines[i]
        stripped = line.strip()

        if stripped.startswith("```"):
            if in_code:
                add_code_block(doc, code_lines)
                code_lines = []
                in_code = False
            else:
                in_code = True
            i += 1
            continue

        if in_code:
            code_lines.append(line)
            i += 1
            continue

        if not stripped:
            i += 1
            continue

        if stripped.startswith("#"):
            level = len(stripped) - len(stripped.lstrip("#"))
            title = clean_inline_md(stripped[level:].strip())
            doc.add_heading(title, level=min(level, 4))
            i += 1
            continue

        if stripped.startswith("|"):
            # Markdown table: header + separator + body rows
            if i + 1 < len(lines) and is_table_separator(lines[i + 1]):
                header = parse_table_line(lines[i])
                rows = []
                i += 2
                while i < len(lines) and lines[i].strip().startswith("|"):
                    rows.append(parse_table_line(lines[i]))
                    i += 1

                table = doc.add_table(rows=1, cols=len(header))
                table.style = "Table Grid"
                hdr_cells = table.rows[0].cells
                for c, value in enumerate(header):
                    hdr_cells[c].text = value

                for row in rows:
                    cells = table.add_row().cells
                    for c, value in enumerate(row):
                        if c < len(cells):
                            cells[c].text = value
                continue

        if stripped.startswith("- "):
            text = clean_inline_md(stripped[2:].strip())
            doc.add_paragraph(text, style="List Bullet")
            i += 1
            continue

        if re.match(r"^\d+\.\s+", stripped):
            text = clean_inline_md(re.sub(r"^\d+\.\s+", "", stripped))
            doc.add_paragraph(text, style="List Number")
            i += 1
            continue

        doc.add_paragraph(clean_inline_md(stripped))
        i += 1

    if in_code and code_lines:
        add_code_block(doc, code_lines)

    try:
        doc.save(OUTPUT_DOCX)
        print(f"Wrote: {OUTPUT_DOCX}")
    except PermissionError:
        doc.save(OUTPUT_DOCX_ALT)
        print(f"Wrote: {OUTPUT_DOCX_ALT} (primary file was locked)")


if __name__ == "__main__":
    main()
